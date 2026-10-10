import json
import pytest
pytest.importorskip("docx")
from pathlib import Path
import sys


def test_body_audit_is_reproducible_and_not_certified():
    root=Path(__file__).resolve().parents[1]
    sys.path.insert(0,str(root/'scripts'))
    from body_page_audit import run
    result=run();stored=json.loads((root/'results/body_page_audit.json').read_text())
    assert result==stored
    assert result['prose_paragraphs']>0
    assert result['total_pdf_pages']>result['estimated_body_text_pages_at_2500_chars']
    assert result['certified_50_body_pages'] is False


def test_first_appendix_boundary_not_alphabetic_order():
    from docx import Document
    from body_page_audit import run
    d=Document();d.add_heading('1. Introduction',level=1);d.add_paragraph('Real body prose.')
    d.add_heading('Appendix B. Comes first',level=1);d.add_paragraph('EXCLUDED APPENDIX PROSE')
    d.add_heading('Appendix A. Comes later',level=1);d.add_paragraph('MORE EXCLUDED PROSE')
    result=run(document=d,pdf_pages=3)
    assert result['prose_words']==3 and result['prose_paragraphs']==1
    assert result['prose_characters']==len('Real body prose.')
