"""Build Times New Roman research papers (docx) with blue page borders,
numbered equations, tables and figures from real result files."""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def _page_border(section, color="1F4E79", size=12):
    sectPr = section._sectPr
    pb = OxmlElement("w:pgBorders"); pb.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "24"); el.set(qn("w:color"), color)
        pb.append(el)
    sectPr.append(pb)


class Paper:
    def __init__(self, title: str, owner_name: str):
        self.doc = Document(); self.eq = 0; self.fig = 0; self.tab = 0
        st = self.doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
            self.doc.styles[s].font.name = "Times New Roman"
            self.doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        _page_border(self.doc.sections[0])
        t = self.doc.add_paragraph(title, style="Title"); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a = self.doc.add_paragraph(owner_name); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.core_properties.author = owner_name
        self.doc.core_properties.last_modified_by = ""
        self.doc.core_properties.comments = ""

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text):
        para = self.doc.add_paragraph(text); para.paragraph_format.space_after = Pt(6)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def equation(self, text: str) -> int:
        self.eq += 1
        para = self.doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(text); r.italic = True
        para.add_run(f"\t\t({self.eq})")
        return self.eq

    def table(self, header: list[str], rows: list[list], caption: str):
        self.tab += 1
        cap = self.doc.add_paragraph(f"Table {self.tab}. {caption}")
        cap.runs[0].bold = True
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i, hdr in enumerate(header):
            t.rows[0].cells[i].text = str(hdr)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
        for row in t.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in para.runs:
                        run.font.name = "Times New Roman"; run.font.size = Pt(9)

    def figure(self, path: str, caption: str, width: float = 6.0):
        self.fig += 1
        self.doc.add_picture(path, width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.doc.add_paragraph(f"Figure {self.fig}. {caption}"); c.runs[0].italic = True

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path: str):
        self.doc.save(path)
        _purge_non_tnr(path)


def _purge_non_tnr(path: str):
    """Replace every Calibri/Calibri Light reference in the docx (styles and
    theme) with Times New Roman so LibreOffice cannot fall back to Carlito."""
    import shutil, zipfile
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ("word/styles.xml", "word/theme/theme1.xml", "word/document.xml"):
                data = data.replace(b"Calibri Light", b"Times New Roman").replace(b"Calibri", b"Times New Roman")
            zout.writestr(item, data)
    shutil.move(tmp, path)
